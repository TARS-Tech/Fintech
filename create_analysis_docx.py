import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = Document()

    # Define Color Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)       # #0F172A Navy
    COLOR_SECONDARY = RGBColor(29, 78, 216)   # #1D4ED8 Royal Blue
    COLOR_TEXT = RGBColor(51, 65, 85)         # #334155 Slate Text
    COLOR_MUTED = RGBColor(100, 116, 139)     # #64748B Muted Gray
    COLOR_BORDER = "CBD5E1"
    HEX_PRIMARY = "0F172A"
    HEX_SECONDARY = "1D4ED8"
    HEX_LIGHT_BG = "F8FAFC"
    HEX_ALT_ROW = "F1F5F9"

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles helper
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = COLOR_TEXT

    # Helpers
    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_paragraph(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Arial'
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_PRIMARY
        r_text = p.add_run(text)
        r_text.font.name = 'Arial'
        r_text.font.italic = italic
        r_text.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Arial'
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_PRIMARY
        r_text = p.add_run(text)
        r_text.font.name = 'Arial'
        r_text.font.color.rgb = COLOR_TEXT
        return p

    def add_callout(text, title=None, border_hex="1D4ED8", fill_hex="F8FAFC"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        tcPr = cell._tc.get_or_add_tcPr()
        borders_elm = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders_elm)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        if title:
            r_title = p.add_run(f"{title}\n")
            r_title.font.name = 'Arial'
            r_title.font.bold = True
            r_title.font.size = Pt(10.5)
            r_title.font.color.rgb = COLOR_SECONDARY
        r_text = p.add_run(text)
        r_text.font.name = 'Arial'
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = COLOR_TEXT
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # COVER PAGE
    # -------------------------------------------------------------
    p_cov_top = doc.add_paragraph()
    p_cov_top.paragraph_format.space_before = Pt(36)
    p_cov_top.paragraph_format.space_after = Pt(12)
    run_org = p_cov_top.add_run("TARS TECHNOLOGIES — CORPORATE REVIEW DOCUMENTATION")
    run_org.font.name = 'Arial'
    run_org.font.size = Pt(11)
    run_org.font.bold = True
    run_org.font.color.rgb = COLOR_SECONDARY

    p_cov_title = doc.add_paragraph()
    p_cov_title.paragraph_format.space_before = Pt(18)
    p_cov_title.paragraph_format.space_after = Pt(12)
    run_title = p_cov_title.add_run("FinPilot Platform:\nDeep Project Analysis Document")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_cov_sub = doc.add_paragraph()
    p_cov_sub.paragraph_format.space_after = Pt(140)
    run_sub = p_cov_sub.add_run("Comprehensive Architectural, Implementation, Technical Security, Product-Market Fit, and Strategic Review")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = COLOR_MUTED

    # Cover Metadata Block Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Name:", "FinPilot (Digital Loan Distribution Platform)"),
        ("Prepared For:", "Corporate Project Review Committee — TARS Technologies"),
        ("Author / Lead:", "Senior Software Architect & Technical Documentation Lead"),
        ("Document Classification:", "Confidential — Internal Corporate Review"),
        ("Date & Version:", "August 2026 | Version 1.0 (Master Release)")
    ]
    for idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.font.name = 'Arial'
        rk.font.bold = True
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = COLOR_PRIMARY
        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.name = 'Arial'
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = COLOR_TEXT

    doc.add_page_break()

    # Setup Header & Footer for main body
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("FinPilot Platform — Project Analysis Document | TARS Technologies")
    hrun.font.name = 'Arial'
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = COLOR_MUTED

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    frun = fp.add_run("CONFIDENTIAL — FOR INTERNAL CORPORATE REVIEW ONLY")
    frun.font.name = 'Arial'
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # SECTION 1: EXECUTIVE ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("1. Executive Analysis")
    add_paragraph("FinPilot is a digital Corporate Direct Selling Agent (DSA) and loan distribution platform designed and engineered at TARS Technologies. The platform is architected to streamline retail credit acquisition by digitizing customer onboarding, identity verification, financial profiling, automated eligibility scoring, and loan discovery.")
    add_paragraph("This Project Analysis Document provides a rigorous, code-grounded architectural evaluation of the FinPilot ecosystem. Based on empirical analysis of the repository (e:\\Fintech), this document evaluates the technical feasibility, data models, security vault controls, operational workflows, and competitive positioning of the system.")
    
    add_callout(
        "FinPilot acts strictly as a digital distribution channel and lead qualification gateway. The platform DOES NOT disburse funds directly or operate as a licensed lender. All credit underwriting and capital disbursement remain the sole responsibility of partner Banks and Non-Banking Financial Companies (NBFCs).",
        title="CORE BUSINESS BOUNDARY STATEMENT",
        border_hex="1D4ED8"
    )

    add_paragraph("Key Analytical Takeaways:")
    add_bullet("The borrower-facing MVP core is highly developed (~55% overall MVP implementation), featuring end-to-end phone OTP login, JWT session management, profile creation, digital Aadhaar/PAN KYC with AES-256-GCM encryption, a deterministic 80-point eligibility scoring engine, and an interactive loan explorer UI.", "Borrower MVP Maturity: ")
    add_bullet("The administrative and operational infrastructure — specifically the Web Admin Panel, backend lead queue management, manual verification tools, and automated lender routing — is documented in specifications but remains unbuilt in the codebase (0% implemented).", "Operational DSA Gap: ")
    add_bullet("Production security hardening is incomplete. Sensitive PII/OTPs are logged in backend console output, Cloudinary KYC document storage uses public URLs without signed access controls, and rate-limiting and JWT refresh endpoints are currently missing.", "Security & Debt Profile: ")

    # -------------------------------------------------------------
    # SECTION 2: PRODUCT DEFINITION
    # -------------------------------------------------------------
    add_heading_1("2. Product Definition")
    add_paragraph("FinPilot is defined as a B2C customer-facing mobile application backed by a Node.js/Express REST API server and a MongoDB document database. Its primary business objective is to capture qualified loan applicants, aggregate their financial profiles and identity documents, execute automated eligibility screening, and package pre-verified loan leads for distribution to lending partners.")
    
    add_heading_2("2.1 Platform Classification")
    add_paragraph("FinPilot is categorized as a Digital Loan Aggregator & Lead Distribution Gateway (Digital DSA). Unlike traditional direct lenders, FinPilot generates revenue through referral commissions earned when a forwarded customer lead is successfully approved and disbursed by a partner Bank or NBFC.")

    add_heading_2("2.2 Core Product Capabilities")
    add_bullet("Allows frictionless registration and authentication via 6-digit mobile OTP, issuing 7-day access and 30-day refresh JWT tokens.", "Digital Authentication: ")
    add_bullet("Captures complete personal details, demographic data, and profile photos stored in Cloudinary (`fintech/profile`).", "Borrower Profile Vault: ")
    add_bullet("Integrates with Sandbox.co.in APIs for Aadhaar OKYC and PAN verification, securing raw identity numbers using AES-256-GCM encryption at rest while exposing only masked identifiers (`XXXX-XXXX-1234`, `XXXXXX1234`).", "Digital KYC Verification: ")
    add_bullet("Evaluates monthly salary, company details, active loan counts, EMI commitments, and CIBIL score against an 80-point deterministic matrix.", "Financial Scoring Engine: ")
    add_bullet("Presents an interactive catalog of 6 loan products (Personal, Home, Business, Education, Vehicle, Gold) with loan amount sliders and mock partner offers.", "Loan Discovery Catalog: ")

    # -------------------------------------------------------------
    # SECTION 3: PROBLEM ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("3. Problem Analysis")
    add_paragraph("Retail loan acquisition in traditional Indian financial markets is severely handicapped by manual, fragmented operational workflows. FinPilot directly addresses four critical friction points:")

    add_bullet("Borrowers must physically visit branches or deal with disparate field agents, submitting paper documents repeatedly for different loan inquiries.", "Fragmented Borrower Journey: ")
    add_bullet("Paper-based Aadhaar and PAN collection creates massive data privacy risks, identity theft vulnerabilities, and compliance failures under India's Digital Personal Data Protection (DPDP) Act.", "Paper-Based KYC Vulnerabilities: ")
    add_bullet("Lenders spend substantial capital processing unqualified applications that fail basic credit score thresholds or income requirements early in underwriting.", "High Acquisition Costs for Lenders: ")
    add_bullet("Borrowers experience complete opaque silence after submitting applications, with zero real-time status visibility.", "Opaque Application Status: ")

    # -------------------------------------------------------------
    # SECTION 4: TARGET USERS & ACTORS
    # -------------------------------------------------------------
    add_heading_1("4. Target Users / Actors")
    add_paragraph("The FinPilot ecosystem defines three distinct operational actors across the loan lifecycle:")

    # Table of Actors
    t_actor = doc.add_table(rows=4, cols=4)
    t_actor.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Actor Role", "Platform Touchpoint", "Primary Objectives", "Implementation Status"]
    for i, h in enumerate(headers):
        cell = t_actor.cell(0, i)
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    actors_data = [
        ("Borrower (Customer)", "React Native Mobile App", "Register, complete KYC, check loan eligibility, explore loan offers, track application", "IMPLEMENTED (UI & APIs)"),
        ("TARS Admin / Operator", "Web Admin Panel (Web Browser)", "Review customer leads, verify KYC documents, execute manual qualification, forward to lenders", "DOCUMENTED ONLY (0% Code)"),
        ("Partner Lenders (Bank/NBFC)", "Lender API / Email / Portal", "Receive qualified leads, conduct credit risk assessment, approve/reject, disburse capital", "PLANNED / NOT IMPLEMENTED")
    ]

    for row_idx, data in enumerate(actors_data, start=1):
        row = t_actor.rows[row_idx]
        bg = HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.0)
            if col_idx == 3:
                r.font.bold = True
                if "IMPLEMENTED" in text:
                    r.font.color.rgb = RGBColor(22, 163, 74)
                elif "DOCUMENTED" in text:
                    r.font.color.rgb = RGBColor(217, 119, 6)
                else:
                    r.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 5: CURRENT VS TARGET SCOPE
    # -------------------------------------------------------------
    add_heading_1("5. Current vs. Target Scope Analysis")
    add_paragraph("A rigorous audit of the codebase reveals a distinct divergence between the current implemented MVP scope and the ultimate target enterprise vision.")

    t_scope = doc.add_table(rows=7, cols=3)
    t_scope.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Functional Area", "Current Implemented Scope (Codebase Reality)", "Target Enterprise Scope (Production Goal)"]
    for i, h in enumerate(s_headers):
        cell = t_scope.cell(0, i)
        set_cell_background(cell, HEX_PRIMARY)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    scope_data = [
        ("Authentication", "Phone + 6-digit OTP login, JWT Access/Refresh tokens, AsyncStorage.", "Biometric auth, OAuth2, SMS gateway integration, hardware SecureStore."),
        ("User Profile", "Personal details, address, profile photo stored in Cloudinary.", "Verified address via PIN API, employment verification, automated geo-tagging."),
        ("KYC Verification", "Aadhaar OKYC via Sandbox.co.in / Mock, PAN verify, AES-256 encryption.", "Direct DigiLocker API, Video KYC (V-KYC), optical character recognition (OCR)."),
        ("Financial Scoring", "Deterministic 80-point score (CIBIL + Salary + Loans) stored in MongoDB.", "Real-time bureau pull (CIBIL/Experian API), Account Aggregator bank statement fetch."),
        ("Loan Offers", "Static local catalog (`loans.js`) with 6 categories & mock partner cards.", "Dynamic real-time lender offer engine, risk-based pricing, direct lender APIs."),
        ("Lead Application", "UI exploration; NO `LoanApplication` model or backend submission API.", "Complete application lifecycle tracking, automated lead routing, lender webhooks.")
    ]

    for row_idx, data in enumerate(scope_data, start=1):
        row = t_scope.rows[row_idx]
        bg = HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.0)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 6: USER JOURNEY ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("6. User Journey Analysis")
    add_paragraph("The mobile application executes a linear customer onboarding flow managed by `AppRouter.jsx` and `RootNavigator.js`. The state transition relies on session token availability and profile completeness.")

    add_heading_2("6.1 Mobile Application Navigation Lifecycle")
    add_bullet("Mobile app boots up, displays SplashScreen, and invokes `StartupScreen.jsx`.", "1. Startup & Hydration: ")
    add_bullet("`StartupScreen` queries `AsyncStorage` for `accessToken` & `refreshToken`. If absent, redirects to `AuthNavigator`.", "2. Token Check: ")
    add_bullet("Customer enters phone and email on `RegisterScreen.jsx`. Server generates 6-digit OTP.", "3. Registration: ")
    add_bullet("Customer submits OTP on `VerificationCode.jsx`. Server validates OTP, returns JWT tokens, and sets Redux `isAuthenticated=true`.", "4. OTP Verification: ")
    add_bullet("Customer is routed to `AppNavigator.jsx`. If `profileCompleted=false`, customer is prompted to complete personal details on `ProfileScreen.jsx`.", "5. Profile Completion: ")
    add_bullet("Customer accesses `KycScreen.jsx`, verifying Aadhaar on `AadhaarScreen.jsx` and PAN on `PanScreen.jsx`.", "6. KYC Verification: ")
    add_bullet("Customer adjusts sliders on `EligibilityScreen.jsx`. Server scores financial data and opens eligibility status modal.", "7. Eligibility Scoring: ")
    add_bullet("Customer browses `LoanExplorerScreen.jsx` and filters partner cards on `LoanOffersScreen.jsx`.", "8. Loan Discovery: ")

    # -------------------------------------------------------------
    # SECTION 7: TECHNICAL ARCHITECTURE ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("7. Technical Architecture Analysis")
    add_paragraph("FinPilot follows a decoupled 3-tier client-server architecture consisting of a React Native mobile client, a stateless Node.js REST API server, and a MongoDB database layer.")

    add_callout(
        "Client (React Native / Expo)  <--->  REST API (Node.js / Express)  <--->  Database (MongoDB / Mongoose)\n"
        "                                           │\n"
        "                                           ├── Cloudinary API (Document Storage)\n"
        "                                           └── Sandbox.co.in API (Aadhaar/PAN Verification)",
        title="HIGH-LEVEL SYSTEM INTERACTION ARCHITECTURE",
        border_hex="0F172A"
    )

    # -------------------------------------------------------------
    # SECTION 8: TECHNOLOGY STACK ASSESSMENT
    # -------------------------------------------------------------
    add_heading_1("8. Technology Stack Assessment")
    add_paragraph("The technology stack has been evaluated for technical suitability, maintainability, and scalability:")

    add_bullet("React Native (Expo SDK ~52.0), Redux Toolkit, Redux Persist, React Navigation v6, React Native Reanimated v3, Axios, Expo Image Picker.", "Frontend Stack: ")
    add_bullet("Node.js (v18+), Express.js, Mongoose ODM, Zod Schema Validation, jsonwebtoken, crypto (built-in Node module), Multer, multer-storage-cloudinary.", "Backend Stack: ")
    add_bullet("MongoDB Atlas / Community Server (Document Database).", "Database Layer: ")
    add_bullet("Sandbox.co.in API (Aadhaar OKYC & PAN match), Cloudinary (Image Cloud Storage).", "Integrations: ")

    # -------------------------------------------------------------
    # SECTION 9: BACKEND ARCHITECTURE ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("9. Backend Architecture Analysis")
    add_paragraph("The backend directory (`backend/src`) is structured using a Layered Controller-Service-Model Pattern, enforcing strict separation of concerns.")

    add_bullet("Defines endpoint mappings, applying authentication and file upload middleware.", "Route Layer (`routes/`): ")
    add_bullet("Executes Zod payload validation (`validations/`), parses HTTP requests, calls business services, and formats JSON responses.", "Controller Layer (`controllers/`): ")
    add_bullet("Encapsulates core business logic (crypto encryption, OTP generation, Sandbox API HTTPS calls, eligibility scoring).", "Service Layer (`services/`): ")
    add_bullet("Defines Mongoose schemas, data types, indexes, and field constraints.", "Model Layer (`models/`): ")

    # -------------------------------------------------------------
    # SECTION 10: MOBILE ARCHITECTURE ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("10. Mobile Architecture Analysis")
    add_paragraph("The mobile repository (`mobile/src`) enforces clean state management using Redux Toolkit combined with Redux Persist for local state caching.")

    add_bullet("Root store (`redux/store.js`) combines `auth`, `user`, `kyc`, and `eligibility` slices. Whitelists `auth` and `kyc` in AsyncStorage.", "Redux Global State: ")
    add_bullet("Centralized API functions (`api/`) wrap Axios calls to backend endpoints.", "API Service Layer: ")
    add_bullet("Custom design system utilizing `Colors` (`theme/colors.js`), custom fonts, safe-area wrappers (`Screen.js`), and shimmer skeleton loaders (`HomeSkeleton.js`).", "UI/UX Components: ")

    # -------------------------------------------------------------
    # SECTION 11: DATABASE & DATA MODEL ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("11. Database & Data Model Analysis")
    add_paragraph("The database layer consists of 5 active Mongoose schemas stored in MongoDB. The schemas utilize unique indexes and object references to maintain data integrity.")

    # Schema Table
    t_model = doc.add_table(rows=6, cols=4)
    t_model.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_headers = ["Model Name", "Collection", "Primary Fields & Types", "Key Constraints & Security"]
    for i, h in enumerate(m_headers):
        cell = t_model.cell(0, i)
        set_cell_background(cell, HEX_PRIMARY)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    models_data = [
        ("User", "users", "phone (Number), email (String), name, dob, profileImage, gender, address, city, state, pincode, isVerified, profileCompleted, kycCompleted, refreshToken", "phone: unique, required; email: unique, sparse; timestamps enabled."),
        ("Otp", "otps", "phone (String), otp (String), expiresAt (Date), attempts (Number)", "phone: index=true; TTL expiration logic in controller."),
        ("KYC", "kycs", "userId (Ref: User), aadhaar (Obj), pan (Obj), aadhaarFront, aadhaarBack, panImage, aadhaarVerified, panVerified, status, kycCompleted", "userId: unique, required; aadhaar/pan encrypted fields (encrypted, iv, authTag, masked)."),
        ("FinancialProfile", "financialprofiles", "userId (Ref: User), monthlySalary, employmentType, company, existingLoans, existingEmi, cibilScore, cibilSource, cibilVerified", "userId: unique, required, index; cibilScore min:300, max:900."),
        ("Eligibility", "eligibilities", "user (Ref: User), score (Number), status (enum: eligible, not_eligible), breakdown (cibil, salary, existingLoans), calculatedAt", "user: unique, required, index; score min:0, max:80.")
    ]

    for row_idx, data in enumerate(models_data, start=1):
        row = t_model.rows[row_idx]
        bg = HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 12: ELIGIBILITY ENGINE ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("12. Loan Eligibility Engine Analysis")
    add_paragraph("The core business differentiator of FinPilot is its automated 80-point eligibility scoring engine (`backend/src/services/eligibility.service.js`). The calculation is deterministic and server-enforced.")

    add_heading_2("12.1 Scoring Matrix Breakdown")
    add_paragraph("The scoring model evaluates three financial parameters:")

    add_bullet("Score >= 750 -> 30 Points | Score >= 700 -> 20 Points | Score >= 650 -> 10 Points | Score < 650 -> 0 Points.", "1. CIBIL Score Points (Max 30): ")
    add_bullet("Salary >= ₹50,000 -> 20 Points | Salary >= ₹30,000 -> 10 Points | Salary < ₹30,000 -> 0 Points.", "2. Monthly Salary Points (Max 20): ")
    add_bullet("0 Active Loans -> 30 Points | 1 Active Loan -> 20 Points | 2 Active Loans -> 10 Points | >= 3 Active Loans -> 0 Points.", "3. Existing Loans Points (Max 30): ")

    add_heading_2("12.2 Decision Threshold Formula")
    add_callout(
        "Total Score = CIBIL Points + Salary Points + Existing Loan Points  (Maximum: 80 Points)\n\n"
        "Eligibility Decision = IF (Total Score >= 70) THEN 'eligible' ELSE 'not_eligible'",
        title="DETERMINISTIC SCORING FORMULA",
        border_hex="1D4ED8"
    )

    add_paragraph("Notice: The eligibility score represents an internal preliminary qualification metric. It DOES NOT guarantee loan approval by partner banks.")

    # -------------------------------------------------------------
    # SECTION 13: KYC & SECURITY ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("13. KYC & Security Architecture Analysis")
    add_paragraph("FinPilot implements robust cryptographic data protection for sensitive customer PII using AES-256-GCM authenticated encryption (`backend/src/services/encryption.service.js`).")

    add_heading_2("13.1 Cryptographic Implementation")
    add_bullet("Utilizes Node.js `crypto` module with AES-256-GCM mode.", "Algorithm: ")
    add_bullet("Requires a 32-byte (256-bit) base64 key configured in `process.env.ENCRYPTION_KEY`.", "Encryption Key: ")
    add_bullet("Generates a random 12-byte Initialization Vector (IV) per encryption operation.", "IV & AuthTag: ")
    add_bullet("Returns an object containing `encrypted`, `iv`, and `authTag` strings, guaranteeing payload confidentiality and tamper-proof authenticity.", "Cipher Storage: ")
    add_bullet("Utility `maskSensitiveData.js` formats Aadhaar as `XXXX-XXXX-1234` and PAN as `XXXXXX1234` for API responses.", "Masking: ")

    # -------------------------------------------------------------
    # SECTION 14: API & INTEGRATION ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("14. API & Integration Analysis")
    add_paragraph("External integrations are managed via backend service layers to prevent exposing API keys to the mobile client:")

    add_bullet("Invokes `https://api.sandbox.co.in/kyc/aadhaar/okyc/otp` and `/verify` using OAuth access tokens generated via `sandbox.service.js`.", "Sandbox.co.in Aadhaar OKYC: ")
    add_bullet("Invokes `https://api.sandbox.co.in/kyc/pan/verify` to validate PAN number, name, and DOB match.", "Sandbox.co.in PAN Verification: ")
    add_bullet("Multer middleware handles direct image multipart uploads to Cloudinary storage buckets (`fintech/profile` and `fintech/kyc`).", "Cloudinary Storage: ")

    # -------------------------------------------------------------
    # SECTION 15: UX/PRODUCT EXPERIENCE ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("15. UX / Product Experience Analysis")
    add_paragraph("The mobile frontend (`mobile/src/screens`) focuses on modern, low-friction financial user experience. Key design highlights include:")

    add_bullet("Dynamic header that collapses smoothly into a compact navbar as the user scrolls (`CollapsibleHeaderLayout.jsx`).", "Collapsible Header Navigation: ")
    add_bullet("Visual multi-variant cards (`LoanCard.jsx`) organizing loan categories into an engaging bento layout.", "Bento Grid Loan Explorer: ")
    add_bullet("Shimmer loaders feedback during network API calls (`HomeSkeleton.js`).", "Skeleton Loaders: ")

    # -------------------------------------------------------------
    # SECTION 16: PRODUCT-MARKET ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("16. Product-Market Analysis")
    add_paragraph("India's digital lending ecosystem is undergoing massive expansion, driven by UPI adoption, India Stack (Aadhaar OKYC, DigiLocker), and growing retail credit demand. FinPilot positions itself at the intersection of retail borrowers and institutional lenders, capturing commission revenue without taking credit balance sheet risk.")

    # -------------------------------------------------------------
    # SECTION 17: MARKET POSITIONING
    # -------------------------------------------------------------
    add_heading_1("17. Market Positioning")
    add_paragraph("FinPilot positions itself as a 'Phygital' Lead Distribution Gateway — combining the high customer trust of local DSA networks with the speed and automation of digital loan aggregators.")

    # -------------------------------------------------------------
    # SECTION 18: MARKET DIFFERENTIATION
    # -------------------------------------------------------------
    add_heading_1("18. Market Differentiation")
    add_paragraph("Unlike legacy aggregators that broadcast user contact details to multiple aggressive telecalling call centers, FinPilot differentiates on:")

    add_bullet("Borrower PII (Aadhaar/PAN) is stored with AES-256-GCM encryption and never sold to third-party telemarketers.", "1. Privacy-First Identity Vault: ")
    add_bullet("Calculates transparent eligibility score before submitting applications to lenders.", "2. Upfront Eligibility Scoring: ")
    add_bullet("Focuses on pre-verified, high-quality lead distribution rather than raw unvetted lead spamming.", "3. Curated Lead Routing: ")

    # -------------------------------------------------------------
    # SECTION 19: COMPETITIVE POSITIONING FRAMEWORK
    # -------------------------------------------------------------
    add_heading_1("19. Competitive Positioning Framework")
    add_paragraph("Formal competitor validation requires external market research. However, based on architecture and product design, FinPilot is evaluated against traditional DSA networks and digital aggregators:")

    t_comp = doc.add_table(rows=4, cols=4)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = ["Dimension", "Traditional DSA Network", "Digital Aggregators (e.g. Paisabazaar)", "FinPilot Platform Approach"]
    for i, h in enumerate(c_headers):
        cell = t_comp.cell(0, i)
        set_cell_background(cell, HEX_PRIMARY)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.0)

    comp_data = [
        ("Onboarding", "Manual paper forms, physical visits", "Web forms, heavy telecalling", "100% Mobile app, self-serve OTP & KYC"),
        ("Data Security", "Physical paper copies, high loss risk", "Unencrypted lead sharing", "AES-256-GCM encrypted database vault"),
        ("Lead Quality", "Unscreened paper files", "Raw unvetted phone leads", "Pre-qualified 80-point eligibility scoring")
    ]

    for row_idx, data in enumerate(comp_data, start=1):
        row = t_comp.rows[row_idx]
        bg = HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTIONS 20-23: SWOT ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("20. Strengths (SWOT Analysis)")
    add_bullet("Fully functional phone OTP, JWT session hydration, profile CRUD, Aadhaar OKYC, PAN verification, and eligibility scoring.", "Solid Borrower Core: ")
    add_bullet("AES-256-GCM encryption for Aadhaar and PAN numbers at rest.", "Strong Encryption: ")
    add_bullet("Zod input schemas on all backend controllers prevent malformed API requests.", "Strict Payload Validation: ")

    add_heading_1("21. Weaknesses (SWOT Analysis)")
    add_bullet("Web Admin Panel and operations lead queue are completely missing in code (0% built).", "Missing Admin Portal: ")
    add_bullet("Loan catalog data (`loans.js`) is hardcoded local JSON rather than database-driven.", "Static Loan Offers: ")
    add_bullet("`LoanApplication` document schema and submission APIs are not yet created.", "No Application Backend: ")

    add_heading_1("22. Opportunities (SWOT Analysis)")
    add_bullet("Integrate DigiLocker and Account Aggregator (AA) APIs for automated financial verification.", "Account Aggregator Framework: ")
    add_bullet("Direct API integrations with partner NBFCs for instant loan disbursal.", "Lender API Integration: ")

    add_heading_1("23. Threats (SWOT Analysis)")
    add_bullet("Compliance enforcement under DPDP Act requires immediate remediation of public Cloudinary document storage.", "Data Protection Regulations: ")
    add_bullet("Competitors offering instant pre-approved credit lines.", "Market Competition: ")

    # -------------------------------------------------------------
    # SECTIONS 24-25: RISKS ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("24. Technical Risks")
    add_bullet("Backend console logs print plaintext OTPs and PII (PAN/Aadhaar) (`auth.controller.js`, `pan.service.js`).", "PII/OTP Logging in Console: ")
    add_bullet("Uploaded KYC document images in Cloudinary use public unauthenticated URLs.", "Public Cloudinary Document Access: ")
    add_bullet("Lack of `/api/auth/refresh` route causes abrupt 401 logouts when 7-day access tokens expire.", "Missing Token Refresh Endpoint: ")
    add_bullet("Mobile client uses plain `AsyncStorage` instead of hardware `SecureStore`.", "Unencrypted Local Token Storage: ")

    add_heading_1("25. Business Risks")
    add_bullet("Without the Web Admin Panel, internal operators cannot view, qualify, or forward collected customer leads.", "Operational Bottleneck: ")
    add_bullet("Showing static mock loan offers risks customer dissatisfaction if actual bank rates differ.", "Mock Rate Discrepancy: ")

    # -------------------------------------------------------------
    # SECTION 26: TESTING / QA ANALYSIS
    # -------------------------------------------------------------
    add_heading_1("26. Testing / QA Analysis")
    add_paragraph("The current codebase relies primarily on manual testing and ad-hoc Node.js scripts (`test_all_scenarios.js`, `test-encryption.js`). There are zero automated Jest unit tests or integration test suites present in either frontend or backend repositories.")

    # -------------------------------------------------------------
    # SECTION 27: PRODUCTION READINESS ASSESSMENT
    # -------------------------------------------------------------
    add_heading_1("27. Production Readiness Assessment")
    add_callout(
        "FinPilot is currently INELIGIBLE FOR PRODUCTION DEPLOYMENT.\n\n"
        "While the borrower mobile MVP functions effectively in local development, production launch is blocked by critical security risks (plaintext PII logging, public Cloudinary URLs, open CORS) and incomplete operational capabilities (missing Admin Panel and Loan Application backend).",
        title="PRODUCTION READINESS: NOT READY",
        border_hex="D97706",
        fill_hex="FFFBEB"
    )

    # -------------------------------------------------------------
    # SECTION 28: STRATEGIC RECOMMENDATIONS
    # -------------------------------------------------------------
    add_heading_1("28. Strategic Recommendations")
    add_paragraph("To transition FinPilot to a production-ready enterprise state, the engineering team must execute the following prioritizations:")

    add_bullet("Remove plaintext console logs, configure private Cloudinary signed URLs, restrict CORS origins, add rate-limiting, and implement JWT token refresh.", "1. Immediate Security Hardening (Priority 1): ")
    add_bullet("Develop `LoanApplication` model, submission controllers, and status tracking APIs.", "2. Loan Application Backend (Priority 2): ")
    add_bullet("Build React Web Admin Panel for lead qualification and operator review.", "3. Admin Web Panel (Priority 3): ")

    # -------------------------------------------------------------
    # SECTION 29: FINAL ASSESSMENT
    # -------------------------------------------------------------
    add_heading_1("29. Final Assessment")
    add_paragraph("FinPilot demonstrates high architectural quality and strong engineering execution across its borrower-facing mobile application and core API layer. The cryptographic implementation (AES-256-GCM), Zod schema validations, and deterministic eligibility scoring provide a resilient technical foundation.")
    add_paragraph("By completing the upcoming security hardening phase and developing the operational Web Admin Panel, FinPilot will successfully transform into a commercial-grade, secure digital loan distribution gateway for TARS Technologies.")

    # Ensure output directory exists
    output_dir = r"e:\Fintech\docs\reports"
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "FinPilot_Project_Analysis_Document.docx")
    
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
